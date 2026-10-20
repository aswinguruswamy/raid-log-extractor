"""Tests for KAN-7: read .txt, .docx, and .vtt files."""

import pytest

from raid.parse import EmptyNotesError, UnsupportedFileError, read_notes

TEAMS_VTT = """WEBVTT

a1b2c3d4-0001
00:00:01.000 --> 00:00:04.500
<v Priya Shah>Morning all. Let's go through the migration plan.</v>

a1b2c3d4-0002
00:00:05.000 --> 00:00:09.000
<v Priya Shah>The vendor may not deliver the API by Friday.</v>

a1b2c3d4-0003
00:00:09.500 --> 00:00:12.000
<v Tom Lee>I will send the test data by Monday.</v>
"""

ZOOM_VTT = """WEBVTT

1
00:00:01.000 --> 00:00:03.000
Priya Shah: We decided to use PostgreSQL.

2
00:00:03.500 --> 00:00:06.000
Tom Lee: Security sign-off is still pending.
"""


def test_txt_gives_plain_text(tmp_path):
    notes = tmp_path / "notes.txt"
    notes.write_text("Risk: the vendor is late.\n\n\n\nAction: Tom sends the plan.", encoding="utf-8")
    assert read_notes(notes) == "Risk: the vendor is late.\n\nAction: Tom sends the plan."


def test_txt_with_windows_encoding(tmp_path):
    notes = tmp_path / "notes.txt"
    notes.write_bytes("Café launch – risk noted".encode("cp1252"))
    assert read_notes(notes) == "Café launch – risk noted"


def test_docx_gives_paragraphs_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Weekly call notes")
    document.add_paragraph("The staging server is down.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Tom Lee"
    path = tmp_path / "notes.docx"
    document.save(path)

    text = read_notes(path)
    assert "Weekly call notes" in text
    assert "The staging server is down." in text
    assert "Owner | Tom Lee" in text


def test_teams_vtt_removes_time_stamps_and_keeps_speakers(tmp_path):
    path = tmp_path / "call.vtt"
    path.write_text(TEAMS_VTT, encoding="utf-8")
    text = read_notes(path)

    assert "-->" not in text
    assert "00:00" not in text
    assert "a1b2c3d4" not in text
    # Two cues from the same speaker merge into one turn.
    assert text.splitlines() == [
        "Priya Shah: Morning all. Let's go through the migration plan. "
        "The vendor may not deliver the API by Friday.",
        "Tom Lee: I will send the test data by Monday.",
    ]


def test_zoom_vtt_keeps_speakers(tmp_path):
    path = tmp_path / "call.vtt"
    path.write_text(ZOOM_VTT, encoding="utf-8")
    assert read_notes(path).splitlines() == [
        "Priya Shah: We decided to use PostgreSQL.",
        "Tom Lee: Security sign-off is still pending.",
    ]


@pytest.mark.parametrize("name", ["notes.pdf", "notes.mp3", "notes"])
def test_other_file_types_give_a_clear_error(tmp_path, name):
    path = tmp_path / name
    path.write_text("text", encoding="utf-8")
    with pytest.raises(UnsupportedFileError, match="Supported file types: .txt, .docx, .vtt"):
        read_notes(path)


def test_empty_file_gives_a_clear_error(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n", encoding="utf-8")
    with pytest.raises(EmptyNotesError):
        read_notes(path)


def test_missing_file_gives_a_clear_error(tmp_path):
    with pytest.raises(FileNotFoundError):
        read_notes(tmp_path / "missing.txt")
