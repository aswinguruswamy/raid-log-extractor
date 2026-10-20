"""Tests for KAN-12 (Excel), KAN-13 (command line), and KAN-14 (cost and time)."""

from openpyxl import load_workbook

from raid.__main__ import main
from raid.cost import COST_TARGET_USD, format_run, run_cost
from raid.export import COLUMNS, SHEET_NAMES, to_excel
from raid.schema import ItemType, RaidItem, RaidLog

from .test_extract import NOTES, FakeClient, item


def sample_log():
    return RaidLog(source_file="call.vtt", items=[
        RaidItem(type="Risk", description="Vendor may be late.",
                 source_quote="The vendor may not deliver the API by Friday.", quote_verified=True),
        RaidItem(type="Action", description="Tom sends the test data.", owner="Tom Lee",
                 due_date="2026-10-26", source_quote="I will send the test data", quote_verified=False),
    ])


# --- KAN-12: Excel ----------------------------------------------------------

def test_one_sheet_for_each_type_plus_summary(tmp_path):
    path = to_excel(sample_log(), tmp_path / "log.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary"] + [SHEET_NAMES[t] for t in ItemType]


def test_header_row_and_open_status(tmp_path):
    workbook = load_workbook(to_excel(sample_log(), tmp_path / "log.xlsx"))
    for name in SHEET_NAMES.values():
        sheet = workbook[name]
        assert [cell.value for cell in sheet[1]] == [header for header, _ in COLUMNS]
        assert sheet.freeze_panes == "A2"
    risks = workbook["Risks"]
    assert risks["B2"].value == "Vendor may be late."
    assert risks["E2"].value == "Open"
    actions = workbook["Actions"]
    assert (actions["C2"].value, actions["D2"].value, actions["E2"].value) == ("Tom Lee", "2026-10-26", "Open")
    assert actions["G2"].value == "Not found"


def test_status_column_has_a_dropdown(tmp_path):
    sheet = load_workbook(to_excel(sample_log(), tmp_path / "log.xlsx"))["Risks"]
    validations = sheet.data_validations.dataValidation
    assert validations and validations[0].formula1 == '"Open,In progress,Closed"'
    assert "E2" in str(validations[0].sqref)


def test_summary_counts(tmp_path):
    sheet = load_workbook(to_excel(sample_log(), tmp_path / "log.xlsx"))["Summary"]
    rows = {row[0]: row[1] for row in sheet.iter_rows(min_row=6, values_only=True)}
    assert rows["Risks"] == 1 and rows["Actions"] == 1 and rows["Decisions"] == 0
    assert rows["Total"] == 2


# --- KAN-13: command line ---------------------------------------------------

def test_command_writes_excel_to_output_folder(tmp_path, capsys):
    notes = tmp_path / "notes.txt"
    notes.write_text(NOTES, encoding="utf-8")
    out = tmp_path / "output"
    code = main(["extract", str(notes), "--out", str(out)], client=FakeClient([item()]))
    assert code == 0
    assert (out / "notes_raid_log.xlsx").is_file()
    printed = capsys.readouterr().out
    assert "Risks         1" in printed
    assert "Saved:" in printed


def test_command_rejects_unsupported_file(tmp_path, capsys):
    notes = tmp_path / "notes.pdf"
    notes.write_text("x", encoding="utf-8")
    assert main(["extract", str(notes)], client=FakeClient([])) == 2
    assert "Supported file types" in capsys.readouterr().err


def test_command_needs_an_api_key(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr("raid.__main__._load_dotenv", lambda: None)
    notes = tmp_path / "notes.txt"
    notes.write_text(NOTES, encoding="utf-8")
    assert main(["extract", str(notes)]) == 2
    assert "ANTHROPIC_API_KEY is not set" in capsys.readouterr().err


# --- KAN-14: cost and time --------------------------------------------------

def test_cost_for_sonnet():
    # 12,000 in x $2/M + 2,000 out x $10/M = $0.024 + $0.020
    assert abs(run_cost("claude-sonnet-5-5", 12_000, 2_000) - 0.044) < 1e-9


def test_one_hour_transcript_estimate_is_under_target():
    assert run_cost("claude-sonnet-5-5", 13_000, 3_000) < COST_TARGET_USD


def test_unknown_model_has_no_cost():
    assert run_cost("some-other-model", 1000, 1000) is None


def test_command_prints_tokens_cost_and_time(tmp_path, capsys):
    notes = tmp_path / "notes.txt"
    notes.write_text(NOTES, encoding="utf-8")
    main(["extract", str(notes), "--out", str(tmp_path)], client=FakeClient([item()]))
    line = capsys.readouterr().out
    assert "Tokens: 1,200 in / 300 out" in line
    assert "Cost: $0.0054" in line  # 1,200 x $2/M + 300 x $10/M
    assert "Time:" in line


def test_format_run_with_unknown_model():
    assert "unknown" in format_run("x", 1, 1, 0.5)
