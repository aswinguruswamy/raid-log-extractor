"""Tests for the web page (KAN-21, KAN-22). No API call is made."""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from raid.export import to_excel_bytes
from raid.schema import RaidItem, RaidLog

APP = str(Path(__file__).parent.parent / "app.py")


def run_app():
    app = AppTest.from_file(APP, default_timeout=30)
    app.run()
    assert not app.exception, app.exception
    return app


def test_demo_mode_shows_results_without_a_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = run_app()
    assert app.radio[0].value.startswith("Demo")
    assert len(app.selectbox[0].options) == 3
    assert app.subheader[1].value.startswith("2. RAID log")
    assert len(app.tabs) == 6
    assert any("Cost:" in c.value for c in app.caption)


def test_every_demo_sample_renders(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = run_app()
    for option in app.selectbox[0].options:
        app.selectbox[0].select(option).run()
        assert not app.exception, option
        total = sum(int(t.label.rsplit("(", 1)[1].rstrip(")")) for t in app.tabs)
        assert app.subheader[1].value == f"2. RAID log · {total} items"


def test_own_notes_mode_asks_for_a_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = run_app()
    app.radio[0].set_value("Your own notes").run()
    assert not app.exception
    assert app.text_input[0].label == "Anthropic API key"
    assert app.button[0].disabled  # no key and no file yet
    assert len(app.tabs) == 0      # no results before a run


def test_upload_is_read_with_its_suffix():
    import importlib.util, sys
    from unittest import mock

    with mock.patch.dict(sys.modules, {"streamlit": mock.MagicMock()}):
        spec = importlib.util.spec_from_file_location("app_module", APP)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    vtt = b"WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v Ana>We decided to ship on Friday.</v>\n"
    assert module.notes_from_upload("call.vtt", vtt) == "Ana: We decided to ship on Friday."


def test_excel_bytes_open_as_a_workbook():
    log = RaidLog(source_file="x.txt", items=[
        RaidItem(type="Risk", description="The vendor may be late.", source_quote="vendor may be late"),
    ])
    workbook = load_workbook(BytesIO(to_excel_bytes(log)))
    assert workbook.sheetnames[0] == "Summary"
    assert workbook["Risks"]["B2"].value == "The vendor may be late."


def test_dollar_signs_show_as_plain_text(monkeypatch):
    """Two "$" signs in Markdown become a maths formula; the page must escape them."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    app = run_app()
    app.radio[0].set_value("Your own notes").run()
    texts = [m.value for m in app.markdown] + [c.value for c in app.caption]
    for text in texts:
        unescaped = text.replace("\\$", "")
        assert "$" not in unescaped, text
    app.radio[0].set_value("Demo (no API key)").run()
    run_line = next(c.value for c in app.caption if "Cost:" in c.value)
    assert "\\$" in run_line


def test_page_and_excel_use_the_same_column_names():
    """BUG-18: the page table and the Excel sheets must name the fields the same way."""
    from raid.export import COLUMNS

    import importlib.util, sys
    from unittest import mock

    with mock.patch.dict(sys.modules, {"streamlit": mock.MagicMock()}):
        spec = importlib.util.spec_from_file_location("app_cols", APP)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    item = RaidItem(type="Risk", description="The vendor may be late.", source_quote="vendor may be late")
    page_columns = list(module.items_table([item]).columns)
    excel_columns = [name for name, _ in COLUMNS if name not in ("#", "Status")]  # the page leaves these out
    assert page_columns == excel_columns
