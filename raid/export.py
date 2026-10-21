"""Write the RAID log to Excel (story S1-07, KAN-12).

The workbook has a Summary sheet, then one sheet for each of the six item types.
Every item starts with status "Open"; the Status column has a dropdown so the
user can change it later.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .schema import ItemType, RaidLog, Status

COLUMNS = [  # (header, width)
    ("#", 5),
    ("Description", 50),
    ("Owner", 18),
    ("Due date", 12),
    ("Status", 12),
    ("Source quote", 60),
    ("Quote found", 14),
]

SHEET_NAMES = {
    ItemType.RISK: "Risks",
    ItemType.ASSUMPTION: "Assumptions",
    ItemType.ISSUE: "Issues",
    ItemType.DEPENDENCY: "Dependencies",
    ItemType.ACTION: "Actions",
    ItemType.DECISION: "Decisions",
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_WRAP = Alignment(wrap_text=True, vertical="top")
_NOT_FOUND_FILL = PatternFill("solid", fgColor="FCE4D6")


def build_workbook(log: RaidLog) -> Workbook:
    """The RAID log as an openpyxl workbook: a Summary sheet, then one sheet per item type."""
    workbook = Workbook()
    _write_summary(workbook.active, log)
    for item_type, items in log.by_type().items():
        _write_type_sheet(workbook.create_sheet(SHEET_NAMES[item_type]), items)
    return workbook


def to_excel(log: RaidLog, path: str | Path) -> Path:
    """Write the log to an .xlsx file and return its path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(log).save(path)
    return path


def to_excel_bytes(log: RaidLog) -> bytes:
    """The .xlsx file as bytes, for a download button."""
    buffer = BytesIO()
    build_workbook(log).save(buffer)
    return buffer.getvalue()


def _write_summary(sheet, log: RaidLog) -> None:
    sheet.title = "Summary"
    sheet.append(["RAID log"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Source file", log.source_file or "-"])
    sheet.append(["Created", date.today().isoformat()])
    sheet.append([])
    sheet.append(["Item type", "Count"])
    for cell in sheet[5]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for item_type, items in log.by_type().items():
        sheet.append([SHEET_NAMES[item_type], len(items)])
    sheet.append(["Total", len(log.items)])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.cell(row=sheet.max_row, column=2).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 40


def _write_type_sheet(sheet, items) -> None:
    sheet.append([header for header, _ in COLUMNS])
    for column, (_, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"

    for number, item in enumerate(items, start=1):
        checked = {True: "Yes", False: "Not found", None: "-"}[item.quote_verified]
        sheet.append([
            number, item.description, item.owner, item.due_date,
            item.status.value, item.source_quote, checked,
        ])
        for cell in sheet[sheet.max_row]:
            cell.alignment = _WRAP
        if item.quote_verified is False:
            sheet.cell(row=sheet.max_row, column=7).fill = _NOT_FOUND_FILL

    last_row = max(sheet.max_row, 2)
    status_list = ",".join(status.value for status in Status)
    dropdown = DataValidation(type="list", formula1=f'"{status_list}"', allow_blank=True)
    sheet.add_data_validation(dropdown)
    dropdown.add(f"E2:E{max(last_row, 200)}")  # room for rows the user adds by hand
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
